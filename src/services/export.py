import io
import json
from typing import List, Optional, Tuple

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from src.models.entities.ingestion import Ingestion
from src.models.entities.pmc_article import PMCArticle
from src.models.entities.pmc_review import PMCReview

# Colours
_FILL_HEADER = PatternFill("solid", fgColor="1F3864")
_FILL_NEW = PatternFill("solid", fgColor="D9EAD3")
_FILL_CHANGED = PatternFill("solid", fgColor="FFF2CC")
_FILL_DECISION = PatternFill("solid", fgColor="EAD1DC")
_FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)
_FONT_DECISION = Font(bold=True, size=11)

_REVIEWER_COLS = ["Decision", "Comment"]


def _fmt_diff(diff: Optional[dict]) -> str:
    """Human-readable summary of a diff: 'field: "old" → "new"' per line."""
    if not diff:
        return ""
    lines = []
    for field, change in diff.items():
        old = str(change.get("old", ""))[:120]
        new = str(change.get("new", ""))[:120]
        lines.append(f"{field}: \"{old}\" → \"{new}\"")
    return "\n".join(lines)


def _changed_fields(diff: Optional[dict]) -> str:
    if not diff:
        return ""
    return ", ".join(diff.keys())


def _build_rows(review_rows: List[Tuple[PMCReview, PMCArticle, Ingestion]]) -> List[dict]:
    out = []
    for review, article, ingestion in review_rows:
        diff = review.diff or {}
        row = {
            "review_id": review.id,
            "epmc_id": review.epmc_id or "",
            "doi": review.doi or (article.doi if article else ""),
            "pm_id": article.pm_id if article else "",
            "review_type": review.review_type,
            "title": article.title if article else "",
            "pub_year": article.pub_year if article else "",
            "language": article.language if article else "",
            "is_open_access": article.is_open_access if article else "",
            "auto_approved": review.auto_approved,
            "auto_approve_reason": review.auto_approve_reason or "",
            "changed_fields": _changed_fields(diff),
            "diff_detail": _fmt_diff(diff),
            "full_diff_json": json.dumps(diff, ensure_ascii=False) if diff else "",
            "ingestion_id": review.ingestion_id,
            "ingested_at": ingestion.ingested_at.strftime("%Y-%m-%d %H:%M UTC") if ingestion and ingestion.ingested_at else "",
            # Reviewer columns — always empty in export
            "Decision": "",
            "Comment": "",
        }
        out.append(row)
    return out


def build_csv(review_rows: List[Tuple]) -> bytes:
    rows = _build_rows(review_rows)
    if not rows:
        rows = [{}]
    df = pd.DataFrame(rows)
    buf = io.BytesIO()
    df.to_csv(buf, index=False, encoding="utf-8-sig")
    return buf.getvalue()


def build_excel(review_rows: List[Tuple]) -> bytes:
    rows = _build_rows(review_rows)
    df = pd.DataFrame(rows) if rows else pd.DataFrame(columns=[
        "review_id", "epmc_id", "doi", "pm_id", "review_type",
        "title", "pub_year", "language", "is_open_access",
        "auto_approved", "auto_approve_reason", "changed_fields", "diff_detail",
        "full_diff_json", "ingestion_id", "ingested_at", "Decision", "Comment",
    ])

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Review Queue")
        ws = writer.sheets["Review Queue"]

        n_cols = len(df.columns)
        n_rows = len(df)

        # ── Header styling ──────────────────────────────────────────────
        for col_idx, col_name in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = _FONT_HEADER
            cell.fill = _FILL_HEADER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # ── Row styling ─────────────────────────────────────────────────
        decision_col_idx = df.columns.get_loc("Decision") + 1
        comment_col_idx = df.columns.get_loc("Comment") + 1

        for row_idx, row_data in enumerate(rows, start=2):
            fill = _FILL_NEW if row_data["review_type"] == "new" else _FILL_CHANGED
            for col_idx in range(1, n_cols + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = fill
                cell.alignment = Alignment(vertical="top", wrap_text=True)

            # Decision + Comment get distinct styling
            for col_idx in (decision_col_idx, comment_col_idx):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.fill = _FILL_DECISION
                cell.font = _FONT_DECISION

        # ── Decision dropdown validation ────────────────────────────────
        dv = DataValidation(
            type="list",
            formula1='"approve,reject"',
            allow_blank=True,
            showDropDown=False,
        )
        dv.sqref = f"{get_column_letter(decision_col_idx)}2:{get_column_letter(decision_col_idx)}{max(n_rows + 1, 2)}"
        ws.add_data_validation(dv)

        # ── Column widths ───────────────────────────────────────────────
        widths = {
            "review_id": 10, "epmc_id": 14, "doi": 22, "pm_id": 14,
            "review_type": 12, "title": 45, "pub_year": 10,
            "language": 10, "is_open_access": 14,
            "auto_approved": 14, "auto_approve_reason": 22,
            "changed_fields": 28, "diff_detail": 60,
            "full_diff_json": 30, "ingestion_id": 13, "ingested_at": 20,
            "Decision": 14, "Comment": 40,
        }
        for col_idx, col_name in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(col_name, 18)

        # ── Freeze header row ───────────────────────────────────────────
        ws.freeze_panes = "A2"

        # ── Summary sheet ───────────────────────────────────────────────
        if rows:
            ingestion_id = rows[0]["ingestion_id"]
            ingested_at = rows[0]["ingested_at"]
            new_count = sum(1 for r in rows if r["review_type"] == "new")
            changed_count = sum(1 for r in rows if r["review_type"] == "changed")

            ws_sum = writer.book.create_sheet("Summary")
            summary_rows = [
                ("Ingestion ID", ingestion_id),
                ("Ingested At", ingested_at),
                ("Total Pending Reviews", len(rows)),
                ("New Articles", new_count),
                ("Changed Articles", changed_count),
                ("", ""),
                ("Instructions", "Fill the 'Decision' column with 'approve' or 'reject'."),
                ("", "Add any notes in the 'Comment' column."),
                ("", "Return the completed file for import."),
            ]
            for r_idx, (label, value) in enumerate(summary_rows, start=1):
                ws_sum.cell(row=r_idx, column=1, value=label).font = Font(bold=True)
                ws_sum.cell(row=r_idx, column=2, value=value)
            ws_sum.column_dimensions["A"].width = 26
            ws_sum.column_dimensions["B"].width = 50

    return buf.getvalue()
